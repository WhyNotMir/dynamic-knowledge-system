"""Shared helpers for tests.

The DKS pipeline is file-driven, so we need a minimal DOCX generator that
produces content with real section structure: headings that set section_path
on extracted elements, and body paragraphs that add up to at least
MIN_CANDIDATE_CHARS (100) so the center detector classifies the section as
its own candidate rather than falling back to orphan clustering.
"""
from __future__ import annotations

import asyncio
import uuid
import base64
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt


def make_docx(path: Path) -> Path:
    """Write a small DOCX with two sections, each with enough text to form a
    structure candidate.

    Layout:
        # Introduction                 (Heading 1)
            <paragraph x 3, total >100 chars>
        # Methods                      (Heading 1)
            <paragraph x 3, total >100 chars>
    """
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "This is the opening section of a synthetic test document. "
        "It introduces the topic and establishes the context for the rest "
        "of the material that follows below."
    )
    doc.add_paragraph(
        "The introduction also lists the goals and the scope of the work, "
        "so that the reader knows what to expect."
    )
    doc.add_paragraph(
        "Finally, the introduction concludes with a short roadmap of "
        "the remaining sections in the document."
    )

    doc.add_heading("Methods", level=1)
    doc.add_paragraph(
        "In the methods section we describe the procedures used to produce "
        "the results reported later. The description is intentionally "
        "generic so that it resembles a real document."
    )
    doc.add_paragraph(
        "Each step is captured in enough detail to exceed the minimum "
        "fragment length threshold used by the center detector."
    )
    doc.add_paragraph(
        "The methods section ends with a short note on limitations."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def make_rich_docx(path: Path) -> Path:
    """Write a DOCX that exercises richer source-driven extraction.

    Includes:
    - heading + formatted paragraph (bold, italic, monospace)
    - quote paragraph
    - bullet list
    - simple table
    """
    doc = Document()
    doc.add_heading("Rich Content", level=1)

    para = doc.add_paragraph()
    para.add_run("This paragraph has ")
    run = para.add_run("bold")
    run.bold = True
    para.add_run(", ")
    run = para.add_run("italic")
    run.italic = True
    para.add_run(", and ")
    run = para.add_run("inline code")
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    para.add_run(" spans.")

    quote = doc.add_paragraph("Quoted material from the source document.")
    for style_name in ("Intense Quote", "Quote"):
        try:
            quote.style = style_name
            break
        except KeyError:
            continue

    item = doc.add_paragraph("First bullet item", style="List Bullet")
    item = doc.add_paragraph("Second bullet item", style="List Bullet")
    item.paragraph_format.left_indent = item.paragraph_format.left_indent

    code = doc.add_paragraph()
    code.add_run("print('hello world')").font.name = "Courier New"

    image_path = path.parent / "tiny.png"
    image_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNgYAAAAAMAASsJTYQAAAAASUVORK5CYII="
    ))
    doc.add_picture(str(image_path))

    doc.add_paragraph("[1] Supplemental footnote from the source.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def make_multitopic_docx(path: Path) -> Path:
    """Write a DOCX with one H1 topic and nested H2/H3 headings."""
    doc = Document()
    doc.add_heading("Platform", level=1)
    doc.add_paragraph(
        "The platform section introduces the broader topic and provides enough "
        "context to comfortably exceed the clustering threshold for a single topic."
    )
    doc.add_heading("Architecture", level=2)
    doc.add_paragraph(
        "Architecture explains how the platform is decomposed into services, "
        "pipelines, and runtime boundaries."
    )
    doc.add_heading("Worker Layer", level=3)
    doc.add_paragraph(
        "The worker layer handles asynchronous execution, background jobs, "
        "and durable retries."
    )
    doc.add_heading("Operations", level=2)
    doc.add_paragraph(
        "Operations focuses on monitoring, deployment ergonomics, and incident response."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def make_linked_docx(path: Path) -> Path:
    """Write a DOCX whose second topic explicitly mentions the first one."""
    doc = Document()

    doc.add_heading("Attention Is All You Need", level=1)
    doc.add_paragraph(
        "Attention Is All You Need introduces the transformer architecture and "
        "establishes the shift away from recurrence in sequence modelling."
    )
    doc.add_paragraph(
        "The article explains how the model uses self-attention, residual "
        "connections, and feed-forward layers to process tokens efficiently."
    )
    doc.add_paragraph(
        "It also frames the broader impact of the transformer on neural "
        "sequence transduction tasks and later model families."
    )

    doc.add_heading("Neural Sequence Transduction Models", level=1)
    doc.add_paragraph(
        "Neural sequence transduction models often build directly on ideas "
        "from Attention Is All You Need when they adopt transformer-based "
        "architectures for translation and other language tasks."
    )
    doc.add_paragraph(
        "This section compares encoder-decoder systems, training regimes, "
        "and the practical advantages of attention-centric model design."
    )
    doc.add_paragraph(
        "It closes by discussing how later systems adapted the transformer "
        "into a more general-purpose family of sequence models."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def make_pdf(path: Path) -> Path:
    """Write a tiny PDF with one obvious heading and one body paragraph."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "PDF Introduction", fontsize=22, fontname="helv")
    page.insert_textbox(
        fitz.Rect(72, 120, 520, 220),
        (
            "This PDF paragraph exists to exercise raw extraction metadata, "
            "including bounding boxes, page dimensions, and font statistics. "
            "It is intentionally long enough to wrap across several lines so "
            "that the body font size becomes the dominant baseline for the page."
        ),
        fontsize=12,
        fontname="Times-Roman",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def make_two_column_pdf(path: Path) -> Path:
    """Write a PDF with a full-width heading and two body columns."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Two Column Title", fontsize=22, fontname="helv")
    page.insert_textbox(
        fitz.Rect(72, 120, 260, 260),
        (
            "Left column paragraph one. "
            "It should be read before the right column content. "
            "The ordering here matters for reconstruction."
        ),
        fontsize=12,
        fontname="Times-Roman",
    )
    page.insert_textbox(
        fitz.Rect(320, 120, 520, 260),
        (
            "Right column paragraph two. "
            "It should appear after the left column paragraph when extracted."
        ),
        fontsize=12,
        fontname="Times-Roman",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def make_repeated_header_footer_pdf(path: Path) -> Path:
    """Write a two-page PDF with repeated header/footer artifacts."""
    doc = fitz.open()
    for index in range(2):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 32), "Conference Header", fontsize=10, fontname="helv")
        page.insert_text((72, 800), f"Page {index + 1}", fontsize=10, fontname="helv")
        page.insert_text((72, 110), f"Body Section {index + 1}", fontsize=20, fontname="helv")
        page.insert_textbox(
            fitz.Rect(72, 150, 520, 280),
            (
                f"This is the main body content for page {index + 1}. "
                "The repeated header and footer should not survive extraction."
            ),
            fontsize=12,
            fontname="Times-Roman",
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def make_hyphenated_pdf(path: Path) -> Path:
    """Write a PDF where one paragraph is split with a hyphen across lines."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Methods", fontsize=20, fontname="helv")
    page.insert_text((72, 130), "The archi-", fontsize=12, fontname="Times-Roman")
    page.insert_text((72, 146), "tecture is designed for reliable extraction.", fontsize=12, fontname="Times-Roman")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def make_front_matter_pdf(path: Path) -> Path:
    """Write a PDF with author metadata before the real body starts."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 60), "Attention Is All You Need", fontsize=22, fontname="helv")
    page.insert_text((72, 96), "ashish@example.com", fontsize=11, fontname="Times-Roman")
    page.insert_text((72, 112), "Google Research", fontsize=11, fontname="Times-Roman")
    page.insert_text((72, 128), "University of Somewhere", fontsize=11, fontname="Times-Roman")
    page.insert_text((72, 174), "Abstract", fontsize=18, fontname="helv")
    page.insert_textbox(
        fitz.Rect(72, 210, 520, 320),
        (
            "This paper describes a transformer model architecture and keeps "
            "the useful abstract content while removing front matter noise."
        ),
        fontsize=12,
        fontname="Times-Roman",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def make_table_pdf(path: Path) -> Path:
    """Write a PDF with a simple ruled 2x2 table detectable by PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Results", fontsize=20, fontname="helv")

    x0, y0 = 72, 130
    col_widths = [160, 160]
    row_heights = [32, 32]
    total_width = sum(col_widths)
    total_height = sum(row_heights)

    # Outer border.
    page.draw_rect(fitz.Rect(x0, y0, x0 + total_width, y0 + total_height), color=(0, 0, 0), width=1)
    # Vertical split.
    page.draw_line(
        fitz.Point(x0 + col_widths[0], y0),
        fitz.Point(x0 + col_widths[0], y0 + total_height),
        color=(0, 0, 0),
        width=1,
    )
    # Horizontal split.
    page.draw_line(
        fitz.Point(x0, y0 + row_heights[0]),
        fitz.Point(x0 + total_width, y0 + row_heights[0]),
        color=(0, 0, 0),
        width=1,
    )

    page.insert_text((x0 + 12, y0 + 20), "Model", fontsize=11, fontname="helv")
    page.insert_text((x0 + col_widths[0] + 12, y0 + 20), "BLEU", fontsize=11, fontname="helv")
    page.insert_text((x0 + 12, y0 + row_heights[0] + 20), "Transformer", fontsize=11, fontname="helv")
    page.insert_text((x0 + col_widths[0] + 12, y0 + row_heights[0] + 20), "28.4", fontsize=11, fontname="helv")

    page.insert_textbox(
        fitz.Rect(72, 240, 520, 320),
        (
            "The evaluation table summarizes the strongest baseline and the "
            "transformer result without duplicating every table cell as prose."
        ),
        fontsize=12,
        fontname="Times-Roman",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def unique_docx(tmp_dir: Path) -> Path:
    return tmp_dir / f"{uuid.uuid4().hex}.docx"


def unique_pdf(tmp_dir: Path) -> Path:
    return tmp_dir / f"{uuid.uuid4().hex}.pdf"


async def confirm_all_candidates(client, project_id: str, proposal_id: str) -> dict:
    """Flip every `proposed` candidate in a proposal to `confirmed`.

    Tests call this as a stand-in for a user clicking every ✓ in the Review
    UI — without it, `/articles/build` now skips every candidate (because
    `proposed` is treated as "not yet reviewed") and returns zero articles.
    """
    r = await client.post(
        f"/projects/{project_id}/structure/proposals/{proposal_id}/confirm-all"
    )
    assert r.status_code == 200, r.text
    return r.json()


async def wait_for_source_done(client, project_id: str, source_id: str, *, timeout: float = 2.0) -> dict:
    """Poll GET /sources/{id} until status == done or timeout.

    Only useful when the ingestion worker is actually running — in unit tests
    we normally drive `run_ingestion` synchronously, so this helper exists
    for UI-flow tests that pretend to observe the FastAPI endpoints.
    """
    deadline = asyncio.get_event_loop().time() + timeout
    while True:
        r = await client.get(f"/projects/{project_id}/sources/{source_id}")
        r.raise_for_status()
        s = r.json()
        if s["status"] in ("done", "failed"):
            return s
        if asyncio.get_event_loop().time() > deadline:
            return s
        await asyncio.sleep(0.05)
