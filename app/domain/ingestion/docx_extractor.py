from __future__ import annotations

import base64
import re
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.domain.ingestion.common import ExtractedElement

HEADING_STYLE_RE = re.compile(r"^Heading\s*(\d+)$", re.IGNORECASE)
MONOSPACE_HINT_RE = re.compile(r"(courier|consolas|menlo|monaco|code)", re.IGNORECASE)


def extract_docx(file_path: str) -> list[ExtractedElement]:
    doc = DocxDocument(file_path)
    elements: list[ExtractedElement] = []
    heading_stack: list[tuple[int, str]] = []
    idx = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            image_runs = extract_inline_images(para)
            for image in image_runs:
                image["section_path"] = " > ".join(h[1] for h in heading_stack)
                image["position_index"] = idx
                elements.append(ExtractedElement(**image))
                idx += 1
            continue

        list_level = get_list_level(para)
        inline_spans = extract_inline_spans(para)
        el_type, h_level = classify_docx_para(para, list_level=list_level)

        if el_type == "heading" and h_level:
            while heading_stack and heading_stack[-1][0] >= h_level:
                heading_stack.pop()
            heading_stack.append((h_level, text))

        section_path = " > ".join(
            h[1] for h in heading_stack if h[0] < (h_level or 99)
        )

        elements.append(
            ExtractedElement(
                content=text,
                element_type=el_type,
                page_number=None,
                section_path=section_path,
                position_index=idx,
                heading_level=h_level,
                list_level=list_level,
                inline_spans=inline_spans or None,
                meta_json=None,
            )
        )
        idx += 1

        for image in extract_inline_images(para):
            image["section_path"] = section_path
            image["position_index"] = idx
            elements.append(ExtractedElement(**image))
            idx += 1

    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells:
                rows.append(cells)
        if rows:
            elements.append(
                ExtractedElement(
                    content="\n".join(" | ".join(cell for cell in row if cell) for row in rows),
                    element_type="table",
                    page_number=None,
                    section_path=" > ".join(h[1] for h in heading_stack),
                    position_index=idx,
                    heading_level=None,
                    list_level=None,
                    inline_spans=None,
                    meta_json={"rows": rows},
                )
            )
            idx += 1

    return elements


def classify_docx_para(
    para: Paragraph,
    *,
    list_level: int | None,
) -> tuple[str, int | None]:
    style_name = para.style.name or ""
    match = HEADING_STYLE_RE.match(style_name)
    if match:
        return "heading", int(match.group(1))
    lowered_style = style_name.lower()
    if "quote" in lowered_style:
        return "quote", None
    if "footnote" in lowered_style or re.match(r"^\[\d+\]\s", para.text.strip()):
        return "footnote", None
    if "code" in lowered_style or paragraph_looks_monospace(para):
        return "code_block", None
    if list_level is not None or "list" in lowered_style or "bullet" in lowered_style:
        return "list_item", None
    if "caption" in lowered_style:
        return "caption", None
    return "paragraph", None


def run_is_monospace(run: Run) -> bool:
    names = [
        getattr(run.font, "name", None),
        getattr(run.style, "name", None) if run.style else None,
    ]
    return any(name and MONOSPACE_HINT_RE.search(name) for name in names)


def paragraph_looks_monospace(para: Paragraph) -> bool:
    text_runs = [run for run in para.runs if run.text]
    if not text_runs:
        return False
    monospace_runs = sum(1 for run in text_runs if run_is_monospace(run))
    return monospace_runs >= max(1, len(text_runs) // 2)


def extract_inline_spans(para: Paragraph) -> list[dict[str, Any]]:
    spans: list[dict[str, Any]] = []
    cursor = 0

    for run in para.runs:
        text = run.text or ""
        if not text:
            continue

        start = cursor
        end = cursor + len(text)
        style_names: list[tuple[str, dict[str, Any] | None]] = []

        if run.bold:
            style_names.append(("bold", None))
        if run.italic:
            style_names.append(("italic", None))
        if run_is_monospace(run):
            style_names.append(("code", None))
        if run.style and (run.style.name or "").lower() == "hyperlink":
            style_names.append(("link", {}))

        for style, data in style_names:
            spans.append(
                {
                    "start": start,
                    "end": end,
                    "style": style,
                    "data": data,
                }
            )

        cursor = end

    return spans


def get_list_level(para: Paragraph) -> int | None:
    p_pr = getattr(para._p, "pPr", None)
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.ilvl is None:
        return None
    try:
        return int(p_pr.numPr.ilvl.val)
    except (TypeError, ValueError):
        return None


def extract_inline_images(para: Paragraph) -> list[dict[str, Any]]:
    images: list[dict[str, Any]] = []
    rels = para.part.related_parts
    for drawing in para._p.xpath(".//w:drawing"):
        blips = drawing.xpath(".//a:blip")
        for blip in blips:
            embed = blip.get(qn("r:embed"))
            if not embed or embed not in rels:
                continue
            part = rels[embed]
            image_bytes = part.blob
            ext = Path(part.filename).suffix.lower() or ".png"
            images.append(
                {
                    "content": "[Image]",
                    "element_type": "image",
                    "page_number": None,
                    "heading_level": None,
                    "list_level": None,
                    "inline_spans": None,
                    "meta_json": {
                        "ext": ext.lstrip("."),
                        "image_base64": base64.b64encode(image_bytes).decode("ascii"),
                    },
                }
            )
    return images
